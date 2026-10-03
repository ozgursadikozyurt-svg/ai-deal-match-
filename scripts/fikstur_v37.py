# Anahtar.ai v3.7 · 1 Ekim 2026
# tests/fixtures/v37_ornek.xlsx üretir (kurgusal veri): python3 scripts/fikstur_v37.py   (gerekirse: pip install openpyxl)
import openpyxl, datetime, os
wb=openpyxl.Workbook(); ws=wb.active; ws.title="1.HAFTA"
for r in [["İSİM-SOYİSİM","KONUM","NİTELİK","MAX BÜTÇE",None],["ALİ VURAL","FENER ","2+1 ASANSÖR ŞART","BÜTÇE DEĞERİNDE","SATILIK"],["ALİ VURAL","ANTALYA GENELİNDE","İMARLI ARSA ","20.000.0000 TL","SATILIK"],["ZEYNEP ŞAHİN","LARA GENELİ","3+1 YADA GENİŞ 2+1 ASANSÖR ŞART","50.000 TL","KİRALIK"],[None]*5]: ws.append(r)
w2=wb.create_sheet("3. HAFTA")
for r in [[None,None,None,None,"İSİM-SOYİSİM","KONUM","NİTELİK","MAX BÜTÇE",None],[None,None,None,None,"GÜL ARSLAN","FENER,ÇAĞLAYAN","2+1 GENİŞ,BAKIMLI",10000000,"SATILIK"]]: w2.append(r)
w3=wb.create_sheet("Portal")
w3.append(["İlan tarihi","Mülk tipi","Mülk türü","İşlem tipi","M2","Fiyat","İlk Fiyat","İlan Başlığı","İlçe","Mahalle","Oda sayısı","Bulunduğu kat","Bina Yaşı","Devren mi","İlan sahibi türü","İlan sahibi","Ofis","İlan Kaynağı","İlan Url"])
w3.append([datetime.date(2026,9,29),"Konut","Daire","Satılık",230,9750000,10500000,"Geniş 5+1 dubleks","Muratpaşa","Muratpaşa Mah","5+1","4. Kat","26-30","Evet","Emlak Ofisi","AYŞE YILMAZ","Örnek Gayrimenkul","Emlakjet","https://www.emlakjet.com/ilan/ornek-19915896"])
w3.append([datetime.date(2026,9,28),"Ticari","Dükkan & Mağaza","Satılık",120,8000000,None,"Cadde üstü dükkan","Muratpaşa","Şirinyalı Mah",None,"Zemin",None,"Evet","Mülk Sahibi","Hasan K.",None,"Sahibinden.com","https://www.sahibinden.com/ilan/ornek/1342746857/detay"])
w3.append([datetime.date(2026,5,1),"Konut","Daire","Satılık",100,4000000,None,"Eski ilan","Muratpaşa","Fener Mah","2+1","Yüksek Giriş","0","Hayır","Emlak Ofisi","Burak Kaya","Deniz Emlak","Hepsi Emlak","https://www.hepsiemlak.com/ornek-555555"])
for c in w3["A"][1:]: c.number_format="dd.mm.yyyy"
os.makedirs("tests/fixtures",exist_ok=True); wb.save("tests/fixtures/v37_ornek.xlsx")